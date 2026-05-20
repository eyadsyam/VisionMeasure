"""
INFS-211 Final Lab Exam Answer Generator
Database Concepts and Design – Chapter 5 (Relational Algebra)
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page margins ─────────────────────────────────────────────────────────────
section = doc.sections[0]
section.top_margin    = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin   = Cm(2.2)
section.right_margin  = Cm(2.2)

# ── Helper functions ──────────────────────────────────────────────────────────
def set_cell_bg(cell, hex_color):
    """Set table-cell background colour."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def set_cell_borders(cell, border_color='1F3864', border_size=6):
    """Add borders to a cell."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'),   'single')
        border.set(qn('w:sz'),    str(border_size))
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), border_color)
        tcBorders.append(border)
    tcPr.append(tcBorders)

def para(text='', bold=False, size=11, color=None, align=WD_ALIGN_PARAGRAPH.LEFT,
         space_before=0, space_after=6, italic=False, font_name='Calibri'):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.bold   = bold
        run.italic = italic
        run.font.name = font_name
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor(*bytes.fromhex(color))
    return p

def add_formula(p_obj, prefix, sigma_or_pi, subscript_text, suffix):
    """Build a run with a subscript for relational algebra notation."""
    if prefix:
        r = p_obj.add_run(prefix)
        r.font.name = 'Cambria Math'; r.font.size = Pt(12)
    # Greek letter + subscript
    r_main = p_obj.add_run(sigma_or_pi)
    r_main.font.name = 'Cambria Math'; r_main.font.size = Pt(13); r_main.bold = True
    r_main.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    # subscript
    r_sub = p_obj.add_run(subscript_text)
    r_sub.font.name = 'Cambria Math'; r_sub.font.size = Pt(9)
    r_sub.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    rPr = r_sub._r.get_or_add_rPr()
    vertAlign = OxmlElement('w:vertAlign')
    vertAlign.set(qn('w:val'), 'subscript')
    rPr.append(vertAlign)
    if suffix:
        r_suf = p_obj.add_run(suffix)
        r_suf.font.name = 'Cambria Math'; r_suf.font.size = Pt(12)

# ══════════════════════════════════════════════════════════════════════════════
# HEADER TABLE  (university branding)
# ══════════════════════════════════════════════════════════════════════════════
hdr = doc.add_table(rows=1, cols=3)
hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr.style = 'Table Grid'

# Left cell – Arabic text
lc = hdr.rows[0].cells[0]
set_cell_bg(lc, '1F3864')
lp = lc.paragraphs[0]
lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
lr = lp.add_run('المملكة العربية السعودية\nوزارة التعليم – جامعة جازان\nكلية الهندسة وعلوم الحاسب')
lr.font.name = 'Arial'; lr.font.size = Pt(9); lr.bold = True
lr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# Centre cell – logo placeholder + title
cc = hdr.rows[0].cells[1]
set_cell_bg(cc, '1F3864')
cp = cc.paragraphs[0]
cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
cr = cp.add_run('JAZAN UNIVERSITY\n')
cr.font.name = 'Calibri'; cr.font.size = Pt(11); cr.bold = True
cr.font.color.rgb = RGBColor(0xFF, 0xD7, 0x00)
cr2 = cp.add_run('College of Engineering & Computer Science')
cr2.font.name = 'Calibri'; cr2.font.size = Pt(8)
cr2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# Right cell – English heading
rc = hdr.rows[0].cells[2]
set_cell_bg(rc, '1F3864')
rp = rc.paragraphs[0]
rp.alignment = WD_ALIGN_PARAGRAPH.LEFT
rr = rp.add_run('Kingdom of Saudi Arabia\nMinistry of Education\nJazan University')
rr.font.name = 'Calibri'; rr.font.size = Pt(9); rr.bold = True
rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# EXAM TITLE BANNER
# ══════════════════════════════════════════════════════════════════════════════
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_before = Pt(4)
title_p.paragraph_format.space_after  = Pt(4)
tr = title_p.add_run('FINAL LAB EXAM – ANSWER SHEET')
tr.font.name = 'Calibri'; tr.font.size = Pt(15); tr.bold = True
tr.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p.paragraph_format.space_after = Pt(2)
sr = sub_p.add_run('Bachelor in Information Technology  |  Database Concepts and Design')
sr.font.name = 'Calibri'; sr.font.size = Pt(11); sr.italic = True
sr.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

sub_p2 = doc.add_paragraph()
sub_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p2.paragraph_format.space_after = Pt(10)
sr2 = sub_p2.add_run('Course Code: INFS-211   |   Academic Year: 2025–2026   |   Second Semester   |   Exam Date: 20/05/2026')
sr2.font.name = 'Calibri'; sr2.font.size = Pt(10)
sr2.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

# ── Divider line ──────────────────────────────────────────────────────────────
div = doc.add_paragraph()
div.paragraph_format.space_before = Pt(0)
div.paragraph_format.space_after  = Pt(0)
pPr  = div._p.get_or_add_pPr()
pBdr = OxmlElement('w:pBdr')
bot  = OxmlElement('w:bottom')
bot.set(qn('w:val'),   'single')
bot.set(qn('w:sz'),    '12')
bot.set(qn('w:space'), '1')
bot.set(qn('w:color'), '1F3864')
pBdr.append(bot); pPr.append(pBdr)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# STUDENT INFO TABLE
# ══════════════════════════════════════════════════════════════════════════════
info = doc.add_table(rows=3, cols=4)
info.alignment = WD_TABLE_ALIGNMENT.CENTER
info.style = 'Table Grid'
info_data = [
    ('Student Name:', '___________________________', 'Student ID:', '_______________'),
    ('Section No:',   '14611',                       'Exam Date:', '20 / 05 / 2026'),
    ('Course Name:',  'Database Concepts and Design', 'Course Code:', 'INFS-211'),
]
for r_idx, row_data in enumerate(info_data):
    cells = info.rows[r_idx].cells
    for c_idx, val in enumerate(row_data):
        cell = cells[c_idx]
        cp   = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run  = cp.add_run(val)
        is_label = (c_idx % 2 == 0)
        run.bold = is_label
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
        if is_label:
            set_cell_bg(cell, 'DCE6F1')
        set_cell_borders(cell, '1F3864', 4)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# SECTION: EMPLOYEE RELATION (given / reference table)
# ══════════════════════════════════════════════════════════════════════════════
sec1 = para('The EMPLOYEE Relation (Reference)', bold=True, size=12, color='1F3864',
            space_before=6, space_after=4)

emp_cols = ['Fname', 'Minit', 'Lname', 'Ssn', 'Bdate', 'Address', 'Sex', 'Salary', 'Super_ssn', 'Dno']
emp_data = [
    ('John',     'B', 'Smith',   '123456789', '09-Jan-1965', '731 Fondren, Houston TX',   'M', '30,000', '333445555', '5'),
    ('Franklin', 'T', 'Wong',    '333445555', '08-Dec-1955', '638 Voss, Houston TX',      'M', '40,000', '888665555', '5'),
    ('Alicia',   'J', 'Zelaya',  '999887777', '19-Jan-1968', '3321 Castle, Spring TX',    'F', '25,000', '987654321', '4'),
    ('Jennifer', 'S', 'Wallace', '987654321', '20-Jun-1941', '291 Berry, Bellaire TX',    'F', '43,000', '888665555', '4'),
    ('Ramesh',   'K', 'Narayan', '666884444', '15-Sep-1962', '975 Fire Oak, Humble TX',   'M', '38,000', '333445555', '5'),
    ('Joyce',    'A', 'English', '453453453', '31-Jul-1972', '5631 Rice, Houston TX',     'F', '25,000', '333445555', '5'),
    ('Ahmad',    'V', 'Jabbar',  '987987987', '29-Mar-1969', '980 Dallas, Houston TX',    'M', '25,000', '987654321', '4'),
    ('James',    'E', 'Borg',    '888665555', '10-Nov-1937', '450 Stone, Houston TX',     'M', '55,000', 'NULL',      '1'),
]

emp_table = doc.add_table(rows=len(emp_data)+1, cols=len(emp_cols))
emp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
emp_table.style = 'Table Grid'

# Header row
for c_idx, col_name in enumerate(emp_cols):
    cell = emp_table.rows[0].cells[c_idx]
    set_cell_bg(cell, '1F3864')
    set_cell_borders(cell, '4472C4', 6)
    cp   = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r    = cp.add_run(col_name)
    r.bold = True; r.font.name = 'Calibri'; r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# Data rows – highlight Dno=4 and Salary>40000
highlight_dno4   = {'2', '3', '6'}   # 0-based: Zelaya=2, Wallace=3, Jabbar=6
highlight_salary  = {'3', '7'}        # Wallace=3, Borg=7
for r_idx, row_data in enumerate(emp_data):
    row_key = str(r_idx)
    bg = 'FFFFFF'
    if row_key in highlight_dno4 and row_key in highlight_salary:
        bg = 'FFE699'   # both conditions (Wallace)
    elif row_key in highlight_dno4:
        bg = 'E2EFDA'   # Dno=4 (green)
    elif row_key in highlight_salary:
        bg = 'FCE4D6'   # Salary>40000 (orange)

    for c_idx, val in enumerate(row_data):
        cell = emp_table.rows[r_idx+1].cells[c_idx]
        set_cell_bg(cell, bg)
        set_cell_borders(cell, '9DC3E6', 4)
        cp   = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r    = cp.add_run(val)
        r.font.name = 'Calibri'; r.font.size = Pt(9)
        if val == 'NULL':
            r.font.color.rgb = RGBColor(0x99, 0x99, 0x99); r.italic = True

legend_p = para('  Legend:  ', bold=True, size=9, space_before=4, space_after=2)
legend_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
runs_data = [
    ('Green rows', 'E2EFDA', ' – Department 4 (Dno = 4)    '),
    ('Orange row', 'FCE4D6', ' – Salary > 40,000    '),
    ('Yellow row', 'FFE699', ' – Both conditions satisfied'),
]
legend_p2 = doc.add_paragraph()
legend_p2.paragraph_format.space_before = Pt(2)
legend_p2.paragraph_format.space_after  = Pt(10)
legend_p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
for label, _, suffix in runs_data:
    r = legend_p2.add_run(f'■ {label}')
    r.font.name = 'Calibri'; r.font.size = Pt(9); r.bold = True
    r = legend_p2.add_run(suffix)
    r.font.name = 'Calibri'; r.font.size = Pt(9)

# ══════════════════════════════════════════════════════════════════════════════
# Q1.1 HEADING
# ══════════════════════════════════════════════════════════════════════════════
def section_banner(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(f'  {text}  ')
    run.bold = True; run.font.name = 'Calibri'; run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    # Shade paragraph background via highlight
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  '1F3864')
    pPr.append(shd)

section_banner('Q1.1 – Relational Algebra Queries on the EMPLOYEE Relation')

# ══════════════════════════════════════════════════════════════════════════════
# Q1.1 (a)
# ══════════════════════════════════════════════════════════════════════════════
def sub_heading(letter, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(4)
    r1 = p.add_run(f'Question 1.1({letter}):  ')
    r1.bold = True; r1.font.name = 'Calibri'; r1.font.size = Pt(11)
    r1.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    r2 = p.add_run(text)
    r2.font.name = 'Calibri'; r2.font.size = Pt(11)

sub_heading('a', 'Retrieve the records from EMPLOYEE whose department is 4.')

# Explanation
para('The SELECT operation (σ) in Relational Algebra filters tuples from a relation that satisfy '
     'a given condition.  We need all employees whose department number (Dno) equals 4.',
     size=10, space_before=2, space_after=6)

# Formula box
f_p = doc.add_paragraph()
f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
f_p.paragraph_format.space_before = Pt(4)
f_p.paragraph_format.space_after  = Pt(4)
add_formula(f_p, 'Relational Algebra Expression:    ', 'σ', 'Dno = 4', '(EMPLOYEE)')

# Formal notation
formal_p = doc.add_paragraph()
formal_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
formal_p.paragraph_format.space_before = Pt(2)
formal_p.paragraph_format.space_after  = Pt(8)
fr = formal_p.add_run('σ(Dno = 4)(EMPLOYEE)')
fr.font.name = 'Courier New'; fr.font.size = Pt(11); fr.bold = True
fr.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

# ── Result table (a) ──────────────────────────────────────────────────────────
para('Result of  σ(Dno = 4)(EMPLOYEE)  — 3 tuples retrieved:', bold=True,
     size=10, color='1F3864', space_before=4, space_after=4)

res_a_data = [
    ('Alicia',   'J', 'Zelaya',  '999887777', '19-Jan-1968', '3321 Castle, Spring TX',   'F', '25,000', '987654321', '4'),
    ('Jennifer', 'S', 'Wallace', '987654321', '20-Jun-1941', '291 Berry, Bellaire TX',   'F', '43,000', '888665555', '4'),
    ('Ahmad',    'V', 'Jabbar',  '987987987', '29-Mar-1969', '980 Dallas, Houston TX',   'M', '25,000', '987654321', '4'),
]
res_a = doc.add_table(rows=len(res_a_data)+1, cols=len(emp_cols))
res_a.alignment = WD_TABLE_ALIGNMENT.CENTER
res_a.style = 'Table Grid'
for c_idx, col_name in enumerate(emp_cols):
    cell = res_a.rows[0].cells[c_idx]
    set_cell_bg(cell, '2E75B6')
    cp = cell.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cp.add_run(col_name); r.bold = True
    r.font.name = 'Calibri'; r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    set_cell_borders(cell, '2E75B6', 6)
for r_idx, row_data in enumerate(res_a_data):
    for c_idx, val in enumerate(row_data):
        cell = res_a.rows[r_idx+1].cells[c_idx]
        set_cell_bg(cell, 'E2EFDA')
        set_cell_borders(cell, '70AD47', 4)
        cp = cell.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(val); r.font.name = 'Calibri'; r.font.size = Pt(9)
        if c_idx == 9:   # Dno column – emphasise
            r.bold = True; r.font.color.rgb = RGBColor(0x37, 0x5E, 0x23)

explanation_a = para(
    'Explanation: The SELECT operation σ(Dno=4) scans every tuple in EMPLOYEE and returns '
    'only those whose Dno attribute equals 4.  Three employees — Zelaya, Wallace, and Jabbar — '
    'satisfy this condition.  The operation does not remove any attributes; all 10 columns '
    'are preserved in the result.',
    size=10, space_before=8, space_after=10, italic=False)

# ══════════════════════════════════════════════════════════════════════════════
# Q1.1 (b)
# ══════════════════════════════════════════════════════════════════════════════
sub_heading('b', 'Select the tuples where Salary > 40,000.')

para('Again we use the SELECT operation (σ).  The condition is now on the Salary attribute: '
     'we want every employee whose salary is strictly greater than 40,000.',
     size=10, space_before=2, space_after=6)

f_p2 = doc.add_paragraph()
f_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
f_p2.paragraph_format.space_before = Pt(4)
f_p2.paragraph_format.space_after  = Pt(4)
add_formula(f_p2, 'Relational Algebra Expression:    ', 'σ', 'Salary > 40000', '(EMPLOYEE)')

formal_p2 = doc.add_paragraph()
formal_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
formal_p2.paragraph_format.space_before = Pt(2)
formal_p2.paragraph_format.space_after  = Pt(8)
fr2 = formal_p2.add_run('σ(Salary > 40000)(EMPLOYEE)')
fr2.font.name = 'Courier New'; fr2.font.size = Pt(11); fr2.bold = True
fr2.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

para('Result of  σ(Salary > 40000)(EMPLOYEE)  — 2 tuples retrieved:', bold=True,
     size=10, color='1F3864', space_before=4, space_after=4)

res_b_data = [
    ('Jennifer', 'S', 'Wallace', '987654321', '20-Jun-1941', '291 Berry, Bellaire TX', 'F', '43,000', '888665555', '4'),
    ('James',    'E', 'Borg',    '888665555', '10-Nov-1937', '450 Stone, Houston TX',  'M', '55,000', 'NULL',      '1'),
]
res_b = doc.add_table(rows=len(res_b_data)+1, cols=len(emp_cols))
res_b.alignment = WD_TABLE_ALIGNMENT.CENTER
res_b.style = 'Table Grid'
for c_idx, col_name in enumerate(emp_cols):
    cell = res_b.rows[0].cells[c_idx]
    set_cell_bg(cell, 'C55A11')
    cp = cell.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cp.add_run(col_name); r.bold = True
    r.font.name = 'Calibri'; r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    set_cell_borders(cell, 'C55A11', 6)
for r_idx, row_data in enumerate(res_b_data):
    for c_idx, val in enumerate(row_data):
        cell = res_b.rows[r_idx+1].cells[c_idx]
        set_cell_bg(cell, 'FCE4D6')
        set_cell_borders(cell, 'ED7D31', 4)
        cp = cell.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(val); r.font.name = 'Calibri'; r.font.size = Pt(9)
        if c_idx == 7:   # Salary column – emphasise
            r.bold = True; r.font.color.rgb = RGBColor(0x84, 0x32, 0x00)
        if val == 'NULL':
            r.font.color.rgb = RGBColor(0x99, 0x99, 0x99); r.italic = True

para(
    'Explanation: The SELECT operation σ(Salary > 40000) scans all tuples in EMPLOYEE and '
    'returns only those whose Salary is strictly greater than 40,000.  Two employees qualify — '
    'Jennifer Wallace (Salary = 43,000) and James Borg (Salary = 55,000).  All 10 columns '
    'are retained; no projection is applied.',
    size=10, space_before=8, space_after=12, italic=False)

# ══════════════════════════════════════════════════════════════════════════════
# SUMMARY TABLE
# ══════════════════════════════════════════════════════════════════════════════
section_banner('Summary — Relational Algebra Expressions')

doc.add_paragraph()
summ = doc.add_table(rows=3, cols=4)
summ.alignment = WD_TABLE_ALIGNMENT.CENTER
summ.style = 'Table Grid'
summ_hdr = ['Part', 'Task Description', 'Relational Algebra Expression', '# Tuples']
for c_idx, h in enumerate(summ_hdr):
    cell = summ.rows[0].cells[c_idx]
    set_cell_bg(cell, '1F3864')
    set_cell_borders(cell, '1F3864', 6)
    cp = cell.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cp.add_run(h); r.bold = True; r.font.name = 'Calibri'; r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

summ_data = [
    ('1.1 (a)', 'Retrieve employees in Department 4',
     'σ(Dno = 4)(EMPLOYEE)', '3'),
    ('1.1 (b)', 'Retrieve employees with Salary > 40,000',
     'σ(Salary > 40000)(EMPLOYEE)', '2'),
]
row_bgs = ['E9F2FB', 'F2F7ED']
for r_idx, (part, desc, expr, count) in enumerate(summ_data):
    cells = summ.rows[r_idx+1].cells
    for c_idx, val in enumerate([part, desc, expr, count]):
        cell = cells[c_idx]
        set_cell_bg(cell, row_bgs[r_idx])
        set_cell_borders(cell, '9DC3E6', 4)
        cp = cell.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(val); r.font.name = 'Calibri'
        r.font.size = Pt(10)
        if c_idx == 2:
            r.font.name = 'Courier New'; r.bold = True
            r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# RELATIONAL ALGEBRA NOTES (theory backing)
# ══════════════════════════════════════════════════════════════════════════════
section_banner('Theoretical Background — SELECT Operation (σ)')
doc.add_paragraph()

theory = doc.add_table(rows=1, cols=2)
theory.alignment = WD_TABLE_ALIGNMENT.CENTER
theory.style = 'Table Grid'

lc2 = theory.rows[0].cells[0]
set_cell_bg(lc2, 'EBF3FB')
set_cell_borders(lc2, '2E75B6', 4)
lp2 = lc2.paragraphs[0]
lp2.paragraph_format.space_before = Pt(4)
lp2.paragraph_format.space_after  = Pt(4)
lr2 = lp2.add_run('Syntax')
lr2.bold = True; lr2.font.name = 'Calibri'; lr2.font.size = Pt(11)
lr2.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
lp3 = lc2.add_paragraph()
lr3 = lp3.add_run('σ<condition>(R)')
lr3.font.name = 'Courier New'; lr3.font.size = Pt(11); lr3.bold = True
lr3.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
lp4 = lc2.add_paragraph()
lr4 = lp4.add_run('where R is a relation and <condition>\nis a Boolean expression on attributes.')
lr4.font.name = 'Calibri'; lr4.font.size = Pt(10)

rc2 = theory.rows[0].cells[1]
set_cell_bg(rc2, 'F9EBE0')
set_cell_borders(rc2, 'ED7D31', 4)
rp2 = rc2.paragraphs[0]
rp2.paragraph_format.space_before = Pt(4)
rr2 = rp2.add_run('Properties')
rr2.bold = True; rr2.font.name = 'Calibri'; rr2.font.size = Pt(11)
rr2.font.color.rgb = RGBColor(0x84, 0x32, 0x00)
props = [
    '• Returns a subset of tuples (horizontal filter).',
    '• All attributes of R are preserved in the result.',
    '• Degree of result = degree of R.',
    '• Cardinality of result ≤ cardinality of R.',
    '• Conditions can use: =, ≠, <, >, ≤, ≥, AND, OR, NOT.',
]
for prop in props:
    pp = rc2.add_paragraph()
    pr = pp.add_run(prop)
    pr.font.name = 'Calibri'; pr.font.size = Pt(10)

doc.add_paragraph()

# ── Footer ────────────────────────────────────────────────────────────────────
footer_p = doc.add_paragraph()
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_p.paragraph_format.space_before = Pt(8)
fr = footer_p.add_run(
    'INFS-211 Database Concepts and Design  |  Second Semester 2025–2026  |  '
    'Jazan University, College of Engineering & Computer Science'
)
fr.font.name = 'Calibri'; fr.font.size = Pt(9); fr.italic = True
fr.font.color.rgb = RGBColor(0x70, 0x70, 0x70)

# ══════════════════════════════════════════════════════════════════════════════
# SAVE
# ══════════════════════════════════════════════════════════════════════════════
out_path = '/home/user/VisionMeasure/INFS211_Exam_Answer/INFS211_FinalLabExam_Answer.docx'
doc.save(out_path)
print(f'Document saved → {out_path}')
